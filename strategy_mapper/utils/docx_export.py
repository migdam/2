from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_docx_report(output_path, document_title, chunks):
    """
    Export analysis results to a Word document.

    chunks = [
        {
            "text": ...,
            "best": ...,
            "score": ...,
            "explanation": ...
        }
    ]
    """
    doc = Document()

    # Add title
    title = doc.add_heading(f"Strategy Mapper Report: {document_title}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")  # Spacer

    # Process each chunk
    for i, ch in enumerate(chunks, start=1):
        doc.add_heading(f"Chunk {i}", level=2)
        doc.add_paragraph(ch["text"])

        # Add best match in bold
        match_para = doc.add_paragraph()
        match_run = match_para.add_run(f"Best match: {ch['best']} ({ch['score']:.3f})")
        match_run.bold = True
        match_run.font.color.rgb = RGBColor(0, 0, 139)  # Dark blue

        # Add explanation
        doc.add_heading("Explanation", level=3)
        doc.add_paragraph(ch["explanation"])

        doc.add_paragraph("")  # Spacer

    doc.save(output_path)
